#!/usr/bin/env python3
"""多级降级的文本抽取：md/txt → docx → pdf（pdftotext → LibreOffice → 内置解析）

用法:
  python extract_text.py --in 简历.pdf --out resume.txt [--meta-out resume_meta.json]

输出:
  --out       纯文本（UTF-8，无 BOM）
  --meta-out  抽取元信息 JSON（method / pages / warnings），可选

降级链（每一步都会打印明示，用户能看到当前走的是哪条路）:
  .pdf : pdftotext → soffice → 内置轻量解析
  .docx: python-docx → soffice → 报错
  其他 : 直接按文本读取
"""
from __future__ import annotations

import glob
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zlib

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import Log, base_parser, check_input, die, read_text, write_json, write_text  # noqa: E402

# ---------------------------------------------------------------- soffice 定位

_SOFFICE_CANDIDATES = [
    os.environ.get("LIBREOFFICE_PATH", ""),
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    "/usr/bin/soffice",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
]


def find_soffice():
    w = shutil.which("soffice") or shutil.which("soffice.exe")
    if w:
        return w
    for c in _SOFFICE_CANDIDATES:
        if c and os.path.exists(c):
            return c
    # DSH 内置 runtime 可能带一份
    for pat in (
        os.path.expanduser(r"~/.dsh/**/soffice.exe"),
        os.path.expanduser(r"~/.dsh/**/LibreOffice/program/soffice.exe"),
    ):
        hits = glob.glob(pat, recursive=True)
        if hits:
            return hits[0]
    return None


def _run(cmd, log: Log) -> bool:
    try:
        p = subprocess.run(cmd, capture_output=True, timeout=180)
        if p.returncode != 0:
            log.warn(f"命令退出码 {p.returncode}: {' '.join(map(str, cmd[:3]))}…")
            return False
        return True
    except Exception as e:  # FileNotFoundError / TimeoutExpired / OSError
        log.warn(f"执行失败: {e}")
        return False


def via_pdftotext(path: str, log: Log):
    exe = shutil.which("pdftotext") or shutil.which("pdftotext.exe")
    if not exe:
        return None
    with tempfile.TemporaryDirectory() as td:
        out = os.path.join(td, "out.txt")
        if not _run([exe, "-enc", "UTF-8", "-layout", path, out], log):
            return None
        if os.path.exists(out):
            return read_text(out)
    return None


def via_soffice(path: str, log: Log):
    exe = find_soffice()
    if not exe:
        log.info("未找到 LibreOffice，跳过该降级路径")
        return None
    with tempfile.TemporaryDirectory() as td:
        src = os.path.join(td, os.path.basename(path))
        shutil.copy2(path, src)
        if not _run(
            [exe, "--headless", "--norestore", "--convert-to", "txt:Text",
             "--outdir", td, src],
            log,
        ):
            return None
        base = os.path.splitext(os.path.basename(src))[0] + ".txt"
        out = os.path.join(td, base)
        if os.path.exists(out):
            return read_text(out)
    return None


def via_docx(path: str, log: Log):
    try:
        import docx  # python-docx
    except ImportError:
        log.info("未安装 python-docx，跳过该路径")
        return None
    try:
        d = docx.Document(path)
    except Exception as e:
        log.warn(f"python-docx 打开失败: {e}")
        return None
    lines = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


# ------------------------------------------------------- 内置轻量 PDF 解析

def _pdf_streams(data: bytes):
    objs = {}
    for m in re.finditer(rb"(\d+)\s+(\d+)\s+obj(.*?)endobj", data, re.S):
        objs[int(m.group(1))] = m.group(3)

    def get_stream(body):
        m = re.search(rb"stream\r?\n", body)
        if not m:
            return None
        raw = body[m.end():body.find(b"endstream", m.end())]
        if b"FlateDecode" in body[:m.start()]:
            try:
                return zlib.decompress(raw)
            except Exception:
                try:
                    return zlib.decompressobj().decompress(raw)
                except Exception:
                    return None
        return raw

    streams = {}
    for num, body in objs.items():
        if b"stream" in body:
            s = get_stream(body)
            if s is not None:
                streams[num] = s
    return objs, streams


def _parse_cmap(s: bytes) -> dict:
    mp = {}
    txt = s.decode("latin-1", errors="replace")
    for m in re.finditer(r"beginbfchar(.*?)endbfchar", txt, re.S):
        for a, b in re.findall(r"<([0-9A-Fa-f]+)>\s*<([0-9A-Fa-f]+)>", m.group(1)):
            try:
                mp[int(a, 16)] = "".join(
                    chr(int(b[i:i + 4], 16)) for i in range(0, len(b), 4)
                )
            except Exception:
                pass
    for m in re.finditer(r"beginbfrange(.*?)endbfrange", txt, re.S):
        for a, b, c, arr in re.findall(
            r"<([0-9A-Fa-f]+)>\s*<([0-9A-Fa-f]+)>\s*(?:<([0-9A-Fa-f]+)>|\[(.*?)\])",
            m.group(1), re.S,
        ):
            lo, hi = int(a, 16), int(b, 16)
            if c:
                dst = int(c, 16)
                for i in range(lo, hi + 1):
                    mp[i] = chr(dst + (i - lo))
            elif arr:
                for k, it in enumerate(re.findall(r"<([0-9A-Fa-f]+)>", arr)):
                    mp[lo + k] = "".join(
                        chr(int(it[i:i + 4], 16)) for i in range(0, len(it), 4)
                    )
    return mp


def via_builtin_pdf(path: str, log: Log):
    """无第三方库时的兜底解析：FlateDecode 解压 + ToUnicode CMap 解码。"""
    with open(path, "rb") as f:
        data = f.read()
    objs, streams = _pdf_streams(data)

    cmaps = {
        num: _parse_cmap(s)
        for num, s in streams.items()
        if b"beginbfchar" in s or b"beginbfrange" in s
    }
    font_to_cmap = {}
    for num, body in objs.items():
        m = re.search(rb"/ToUnicode\s+(\d+)\s+0\s+R", body)
        if m:
            font_to_cmap[num] = cmaps.get(int(m.group(1)), {})

    name_to_font = {}
    for num, body in objs.items():
        fm = re.search(rb"/Font\s*<<(.*?)>>", body, re.S)
        if not fm:
            continue
        d = {}
        for m in re.finditer(rb"/([A-Za-z0-9#+._-]+)\s+(\d+)\s+0\s+R", fm.group(1)):
            d[m.group(1).decode()] = int(m.group(2))
        if d:
            name_to_font[num] = d

    pages = sorted(
        num for num, body in objs.items()
        if b"/Type /Page" in body and b"/Type /Pages" not in body
    )

    def decode_hex(hexs, cm):
        codes = [int(hexs[i:i + 4], 16) for i in range(0, len(hexs) - 3, 4)]
        return "".join(cm.get(c, "") for c in codes)

    page_texts = []
    for pnum in pages:
        body = objs[pnum]
        rm = re.search(rb"/Resources\s+(\d+)\s+0\s+R", body)
        fonts = {}
        if rm and int(rm.group(1)) in name_to_font:
            fonts = name_to_font[int(rm.group(1))]
        else:
            fm = re.search(rb"/Font\s*<<(.*?)>>", body, re.S)
            if fm:
                for m in re.finditer(rb"/([A-Za-z0-9#+._-]+)\s+(\d+)\s+0\s+R", fm.group(1)):
                    fonts[m.group(1).decode()] = int(m.group(2))

        ccont = re.search(rb"/Contents\s+(\d+)\s+0\s+R", body)
        if ccont:
            refs = [int(ccont.group(1))]
        else:
            m2 = re.search(rb"/Contents\s*\[(.*?)\]", body, re.S)
            refs = [int(x) for x in re.findall(rb"(\d+)\s+0\s+R", m2.group(1))] if m2 else []
        s = b"\n".join(streams[r] for r in refs if r in streams)
        if not s:
            continue

        txt = s.decode("latin-1", errors="replace")
        cur, line, lines = None, [], []
        for m in re.finditer(
            r"/([A-Za-z0-9#+._-]+)\s+[\d.]+\s+Tf|<([0-9A-Fa-f]*)>\s*Tj"
            r"|\[(.*?)\]\s*TJ|\((.*?)\)\s*Tj|(T\*|ET|BT)",
            txt, re.S,
        ):
            if m.group(1):
                cur = m.group(1)
                continue
            cmap = font_to_cmap.get(fonts.get(cur, -1), {})
            if m.group(2) is not None:
                line.append(decode_hex(m.group(2), cmap))
            elif m.group(3) is not None:
                for hm in re.finditer(r"<([0-9A-Fa-f]*)>|\(((?:\\.|[^)])*)\)", m.group(3), re.S):
                    line.append(decode_hex(hm.group(1), cmap) if hm.group(1) is not None else hm.group(2))
            elif m.group(4) is not None:
                line.append(m.group(4))
            elif line:
                lines.append("".join(line))
                line = []
        if line:
            lines.append("".join(line))
        page_texts.append("\n".join(lines))

    if not page_texts:
        return None
    log.info(f"内置解析：{len(page_texts)} 页")
    return {"text": "\n".join(page_texts), "pages": len(page_texts)}


# ------------------------------------------------------------------ 主入口

def extract(path: str, log: Log) -> dict:
    ext = os.path.splitext(path)[1].lower()
    warnings = []

    if ext in (".md", ".markdown", ".txt", ".text", ""):
        return {"text": read_text(path), "method": "plain-text", "pages": None,
                "warnings": warnings}

    if ext == ".docx":
        t = via_docx(path, log)
        if t and t.strip():
            return {"text": t, "method": "python-docx", "pages": None, "warnings": warnings}
        log.warn("python-docx 路径失败，降级 LibreOffice")
        t = via_soffice(path, log)
        if t and t.strip():
            return {"text": t, "method": "libreoffice", "pages": None, "warnings": warnings}
        die("DOCX 解析失败：python-docx 与 LibreOffice 均不可用。请另存为 .md/.txt 后重试")

    if ext == ".pdf":
        for name, fn in (("pdftotext", via_pdftotext), ("libreoffice", via_soffice)):
            log.step(f"尝试 {name}")
            t = fn(path, log)
            if t and t.strip():
                return {"text": t, "method": name, "pages": None, "warnings": warnings}
        log.step("尝试内置轻量解析（兜底）")
        r = via_builtin_pdf(path, log)
        if r and r["text"].strip():
            warnings.append("使用内置解析，复杂排版可能丢失结构；建议对照原文核对")
            return {"text": r["text"], "method": "builtin", "pages": r["pages"],
                    "warnings": warnings}
        die("PDF 解析失败：所有降级路径均失败。请手动复制简历文本粘贴为 .md 后重试")

    # 未知扩展名：当纯文本试一次
    log.warn(f"未知扩展名 {ext}，按纯文本尝试")
    return {"text": read_text(path), "method": "unknown-as-text", "pages": None,
            "warnings": ["未知文件类型，按纯文本解析"]}


def main():
    ap = base_parser("多级降级文本抽取")
    ap.add_argument("--meta-out", dest="meta_out", default=None,
                    help="抽取元信息 JSON 输出路径（可选）")
    # --out 在本脚本是纯文本；为兼容 base_parser 的必填 --out，保留同名参数
    args = ap.parse_args()

    check_input(args.inp)
    log = Log(args.quiet)
    log.step(f"抽取文本: {args.inp}")

    if args.dry_run:
        print(f"[dry-run] 将解析 {args.inp} → 文本 {args.out}"
              + (f" + 元信息 {args.meta_out}" if args.meta_out else ""))
        return

    r = extract(args.inp, log)
    write_text(args.out, r["text"])
    log.info(f"方法={r['method']} 字符数={len(r['text'])} → {args.out}")
    if args.meta_out:
        write_json(args.meta_out, {
            "source": os.path.abspath(args.inp),
            "method": r["method"],
            "pages": r["pages"],
            "chars": len(r["text"]),
            "warnings": r["warnings"],
        })


if __name__ == "__main__":
    main()
